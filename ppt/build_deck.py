"""生成比赛 PPT：python ppt/build_deck.py -> ppt/票据不说谎.pptx

所有数字都来自仓库 README 里记录的实验结果；改数字请先改 README，再改这里。
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "票据不说谎.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
INK = RGBColor(0x1C, 0x24, 0x33)
MUTED = RGBColor(0x6B, 0x73, 0x85)
KICK = RGBColor(0x5B, 0x6B, 0x8C)
RED = RGBColor(0xC7, 0x00, 0x0B)
RED_BG = RGBColor(0xFD, 0xEC, 0xEC)
GREEN = RGBColor(0x17, 0x80, 0x4A)
GREEN_BG = RGBColor(0xE8, 0xF5, 0xEE)
GRAY1 = RGBColor(0x8C, 0x95, 0xA6)
GRAY2 = RGBColor(0xBC, 0xC3, 0xCF)
TINT = RGBColor(0xF3, 0xF5, 0xF9)
LINE = RGBColor(0xE1, 0xE5, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CN = "Microsoft YaHei"
EN = "Arial"
MONO = "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


# ---------------------------------------------------------------- 基础件

def _fonts(run, face, latin=None):
    rPr = run._r.get_or_add_rPr()
    for tag, f in (("a:latin", latin or face), ("a:ea", face), ("a:cs", face)):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", f)


def R(text, size=16, color=INK, bold=False, font=CN, latin=None, italic=False):
    return dict(text=text, size=size, color=color, bold=bold, font=font, latin=latin, italic=italic)


def text(slide, x, y, w, h, paras, anchor="top", margin=0.0, align=PP_ALIGN.LEFT, after=0, line=None):
    """paras：每段是一个 run 列表，或 (run 列表, 段落选项) 二元组。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for i, p in enumerate(paras):
        runs, opt = (p if isinstance(p, tuple) else (p, {}))
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = opt.get("align", align)
        para.space_after = Pt(opt.get("after", after))
        if opt.get("line", line):
            para.line_spacing = opt.get("line", line)
        for r in (runs if isinstance(runs, list) else [runs]):
            run = para.add_run()
            run.text = r["text"]
            run.font.size = Pt(r["size"])
            run.font.bold = r["bold"]
            run.font.italic = r["italic"]
            run.font.color.rgb = r["color"]
            _fonts(run, r["font"], r["latin"])
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def circle_num(slide, x, y, n, d=0.46, fill=NAVY, color=WHITE, size=15):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(n)
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), True, color
    _fonts(run, CN, EN)
    return c


def arrow(slide, x, y, w=0.34, h=0.34, fill=GRAY2):
    return box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.RIGHT_ARROW)


def image(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def header(slide, kicker, title, n):
    text(slide, 0.6, 0.42, 10, 0.3, [[R(kicker, 13, KICK, True)]])
    text(slide, 0.6, 0.72, 12.1, 0.8, [[R(title, 30, NAVY, True)]])
    text(slide, W - 1.2, 7.05, 0.6, 0.25, [[R(str(n), 11, MUTED, latin=EN)]], align=PP_ALIGN.RIGHT)


def footnote(slide, s):
    text(slide, 0.6, 7.02, 11.2, 0.3, [[R(s, 10, MUTED)]])


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def new(dark=False):
    s = prs.slides.add_slide(BLANK)
    if dark:
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = NAVY
    return s


def table(slide, x, y, w, rows, col_w, header_fill=NAVY, size=14, row_h=0.46, fills=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                 Inches(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, cell_v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (fills[i][j] if fills else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            spec = cell_v if isinstance(cell_v, dict) else R(str(cell_v), size, WHITE if i == 0 else INK, i == 0)
            run = p.add_run()
            run.text = spec["text"]
            run.font.size, run.font.bold, run.font.color.rgb = Pt(spec["size"]), spec["bold"], spec["color"]
            _fonts(run, spec["font"], spec["latin"])
    return tbl


def chart_style(chart, legend=True):
    chart.font.size = Pt(12)
    chart.font.color.rgb = INK
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
    va = chart.value_axis
    va.minimum_scale, va.maximum_scale = 0, 1.12     # 顶部留出放数值标签的空间
    va.major_unit = 0.25
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    va.tick_labels.number_format = "0%"
    va.tick_labels.number_format_is_linked = False
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.font.color.rgb = MUTED
    ca = chart.category_axis
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(13)
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = "0.0%"
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12)
    dl.font.color.rgb = INK
    return plot


# ---------------------------------------------------------------- 1 封面

s = new(dark=True)
text(s, 0.8, 0.8, 11, 0.4, [[R("第十七届“工行杯”全国大学生金融科技创新大赛 · 金融安全服务方向", 15, ICE)]])
text(s, 0.8, 1.7, 11, 1.3, [[R("票据不说谎", 60, WHITE, True)]])
text(s, 0.8, 3.0, 11.2, 0.9, [[R("让大模型只负责转写、由代码核对算术的单据可信度核验", 24, ICE)]])
stats = [("96.7%", "篡改检出率", "CORD 公开数据 150 张被改单据"),
         ("3.3%", "误报率", "150 张未改单据"),
         ("p = 0.021", "显著优于让模型自己核算", "预先登记的确认性检验")]
for i, (big, label, sub) in enumerate(stats):
    x = 0.8 + i * 3.95
    box(s, x, 4.35, 3.65, 1.65, fill=RGBColor(0x2A, 0x38, 0x5A))
    text(s, x + 0.3, 4.5, 3.2, 0.7, [[R(big, 34, WHITE, True, latin=EN)]])
    text(s, x + 0.3, 5.2, 3.2, 0.35, [[R(label, 15, WHITE, True)]])
    text(s, x + 0.3, 5.55, 3.2, 0.35, [[R(sub, 12, ICE)]])
text(s, 0.8, 6.6, 11, 0.4, [[R("参赛者：［姓名］　｜　［学校］", 14, ICE)]])
notes(s, "开场一句话：票据是资金流出的依据，我们做的是让票据自己的算术替它作证。"
         "三个数字先亮出来：公开数据上 150 张被改的单据检出 96.7%，150 张未改的误报 3.3%，"
         "而且在预先登记的检验里显著优于“让大模型自己核算”这种直觉做法。")

# ---------------------------------------------------------------- 2 痛点

s = new()
header(s, "问题", "单据是资金流出的依据，核验却还停在“识别”", 2)
cards = [("1", "改金额是最省事的造假", "报销、贸易融资、供应链金融都以单据为付款依据。",
          "1 个数字", "改一处就能多报一笔钱，不必伪造整张单据"),
         ("2", "OCR 只识别、不核验", "读错了、被改了，输出的数字照样格式完好、数量级合理。",
          "47,400", "票面印的是 47,600——读错了，却看不出来"),
         ("3", "直接问大模型也靠不住", "公开数据实测：直接问大模型“这张单据是否被篡改”。",
          "74.7%", "150 张被改的单据里只检出这么多")]
for i, (n, title, body, big, cap) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 1.85, 3.8, 4.6)
    circle_num(s, x + 0.35, 2.15, n)
    text(s, x + 0.35, 2.85, 3.2, 0.9, [[R(title, 20, NAVY, True)]])
    text(s, x + 0.35, 3.55, 3.15, 1.1, [[R(body, 15, INK)]], line=1.2)
    text(s, x + 0.35, 4.75, 3.2, 0.85, [[R(big, 40, RED, True, latin=EN)]])
    text(s, x + 0.35, 5.6, 3.2, 0.6, [[R(cap, 13, MUTED)]], line=1.1)
footnote(s, "47,400：CORD-v2 test #20 的一次真实误读（见下页）。74.7%：CORD-v2 train 300 张（150 张篡改），同一视觉模型，每张读 3 次取多数。")
notes(s, "痛点三层：造假成本低；传统 OCR 只读不核；而最直觉的“直接问大模型”在公开数据上只检出 74.7%。"
         "后面会解释为什么模型自己判断不可靠。")

# ---------------------------------------------------------------- 3 发现

s = new()
header(s, "发现", "大模型会自己把账做平", 3)
image(s, ASSETS / "cord20_totals.png", 0.6, 1.75, w=5.2)
text(s, 0.6, 5.1, 5.2, 0.9, [[R("票面：应付 TOTAL 印的是 377,859，下方没有现金与找零行。", 13, MUTED)]])
rows = [["项目", "票面印的", "模型第 1 次报的"],
        ["米饭一行", "47,600", R("47,400", 15, RED, True, latin=EN)],
        ["小计", "325,600", R("325,400", 15, RED, True, latin=EN)],
        ["应付 TOTAL", "377,859", R("377,659", 15, RED, True, latin=EN)],
        ["现金 CASH", R("（没有这一行）", 14, MUTED), R("400,000", 15, RED, True, latin=EN)],
        ["找零 CHANGE", R("（没有这一行）", 14, MUTED), R("22,341", 15, RED, True, latin=EN)]]
table(s, 6.3, 1.8, 6.4, rows, [2.0, 2.1, 2.3], size=15, row_h=0.5)
box(s, 6.3, 5.0, 6.4, 1.3, fill=RED_BG)
text(s, 6.55, 5.12, 6.0, 1.1, [
    [R("五个数彼此完全自洽，三条校验全部通过——但全是错的。", 16, RED, True)],
    [R("一处误读之后，模型改写了应付、还补出两行票面上不存在的付款，把账做平了。", 14, INK)]],
    line=1.15, after=4)
footnote(s, "CORD-v2 test #20，未经篡改。同一张图另外 2 次识别如实报出了 377,859。")
notes(s, "这是整个设计的起点。这张票没被人改过，模型把 6 读成 4 之后，没有停下来，"
         "而是把应付改成与误读一致的数，还编出了现金和找零两行。结果是一份完全自洽、但全错的读数。"
         "所以不能相信“模型读出来是平的”，更不能让模型自己去判断平不平。")

# ---------------------------------------------------------------- 4 设计

s = new()
header(s, "设计", "铁律：模型只转写，所有算术在代码侧完成", 4)
steps = [("票据图像", "小票、发票\n拍照或扫描"),
         ("视觉大模型", "逐格照抄票面\n不做任何计算"),
         ("代码核对", "票面自带的\n多条等式"),
         ("三次识别", "各自判定\n按多数定结论")]
for i, (t, sub) in enumerate(steps):
    x = 0.6 + i * 2.72
    box(s, x, 2.0, 2.3, 1.9, fill=NAVY if i in (1, 2) else TINT)
    col = WHITE if i in (1, 2) else NAVY
    text(s, x + 0.2, 2.2, 1.9, 0.5, [[R(t, 19, col, True)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 2.8, 1.9, 0.9, [[R(sub, 14, ICE if i in (1, 2) else INK)]], align=PP_ALIGN.CENTER, line=1.15)
    arrow(s, x + 2.33, 2.78)
box(s, 11.48, 2.0, 1.25, 1.9, fill=TINT)
text(s, 11.5, 2.15, 1.2, 1.7, [[R("可信", 16, GREEN, True)], [R("存疑", 16, RED, True)],
                               [R("无法核验", 13, MUTED, True)]], align=PP_ALIGN.CENTER, after=6)
box(s, 0.6, 4.4, 12.13, 1.05, fill=RED_BG)
text(s, 0.9, 4.52, 11.6, 0.85, [[R("提示词里没有任何等式。", 18, RED, True),
                                  R("模型不知道自己被按什么标准检查，也就无法朝那个标准编数字。", 18, INK)]],
     anchor="middle")
text(s, 0.6, 5.75, 12.1, 1.0, [
    [R("每个“存疑”都能指到具体等式：", 15, NAVY, True),
     R("“小计 17,272 + 税 1,427 = 18,699；票面应付 18,999；差 +300”。审核员看的是一条算式，不是一个黑箱分数。", 15, INK)]],
    line=1.2)
notes(s, "设计只有一条铁律：模型只做它擅长的——把票面照抄下来；加减乘除、判断平不平，全部交给代码。"
         "提示词里不出现任何等式，模型就没有“朝答案凑数”的动机。输出不是一个分数，而是失败的那条算式。")

# ---------------------------------------------------------------- 5 冗余

s = new()
header(s, "原理", "票面自带的冗余，就是免费的校验位", 5)
cols = [("零售小票", "港式超市小票、CORD 印尼小票", [
            "商品行合计 − 小计上方的折扣 = 小计",
            "应付 = 小计 − 下方折扣 + 税 + 服务费 + 舍入",
            "付款 − 找零 = 应付",
            "折扣标签自带的金额 = 金额栏（港式）"]),
        ("B2B 商业发票", "贸易融资、对公报销", [
            "每行：数量 × 单价 = 金额",
            "每行：金额 × (1 + 税率) = 含税金额",
            "各行金额之和 = 不含税合计",
            "不含税 + 税额 = 含税合计",
            "不含税 × 税率 = 税额",
            "分税率小计行 = Total 行（合计印了两遍）"])]
for i, (t, sub, eqs) in enumerate(cols):
    x = 0.6 + i * 6.2
    box(s, x, 1.75, 5.9, 3.0)
    text(s, x + 0.35, 1.95, 5.3, 0.45, [[R(t, 20, NAVY, True), R(f"　{sub}", 13, MUTED)]])
    text(s, x + 0.35, 2.55, 5.3, 2.1, [([R("•  " + e, 14.5, INK)], {"after": 5}) for e in eqs])
text(s, 0.6, 5.0, 12.1, 0.4, [[R("想把应付金额改掉又不被发现，至少要同时改：", 15, NAVY, True)]])
stats = [("3–4 个数", "小票：商品行、小计、应付，有付款行时再加付款或找零"),
         ("9 个数", "发票：单价或数量、行金额、行含税、Total 行三个、小计行三个")]
for i, (big, label) in enumerate(stats):
    x = 0.6 + i * 6.2
    text(s, x, 5.5, 2.3, 0.9, [[R(big, 34, RED, True)]], anchor="middle")
    text(s, x + 2.4, 5.5, 3.5, 0.9, [[R(label, 14, INK)]], anchor="middle", line=1.15)
footnote(s, "改一个数字只会破坏 1–3 条等式；但要让所有等式重新闭合，就得把整条依赖链上的数一起改掉。")
notes(s, "为什么算术能抓造假：票据为了让人看得懂，本来就把同一笔钱记了好几遍。"
         "这些冗余对人是多余的，对机器是免费的校验位。单据越规范，冗余越多，造假要同时改的数字越多。")

# ---------------------------------------------------------------- 6 评测设计

s = new()
header(s, "评测", "先调规则、后冻结，每批新数据只跑一次", 6)
stages = [("开发集", "CORD test 前 20 张", "找出并修正 3 处规则问题"),
          ("留出集", "CORD test 33 张", "首次冻结检验：检出 17/17"),
          ("预先登记", "CORD validation 53 张", "加入两个对照组"),
          ("换场景", "商业发票 71 张", "检出 36/36，误报 0/35"),
          ("确认性检验", "CORD train 300 张", "预先登记主检验")]
for i, (t, data, what) in enumerate(stages):
    x = 0.6 + i * 2.47
    circle_num(s, x + 0.875, 1.85, i + 1, d=0.5, fill=RED if i == 4 else NAVY)
    if i < 4:
        box(s, x + 1.42, 2.07, 1.97, 0.06, fill=GRAY2, shape=MSO_SHAPE.RECTANGLE)
    box(s, x, 2.6, 2.25, 1.95, fill=RED_BG if i == 4 else TINT)
    text(s, x + 0.2, 2.75, 1.9, 0.45, [[R(t, 17, RED if i == 4 else NAVY, True)]])
    text(s, x + 0.2, 3.25, 2.0, 0.4, [[R(data, 12, MUTED)]])
    text(s, x + 0.2, 3.7, 1.9, 0.8, [[R(what, 14, INK)]], line=1.15)
box(s, 0.6, 4.95, 12.13, 1.45)
text(s, 0.9, 5.08, 11.6, 1.25, [
    ([R("两个对照组", 15, NAVY, True), R("　同一批图片、同一模型、各读 3 次、同一投票规则，只换“判断由谁来做”", 14, MUTED)], {"after": 6}),
    ([R("对照 A　", 15, INK, True), R("直接问模型：这张单据是否被篡改？", 15, INK)], {"after": 4}),
    [R("对照 B　", 15, INK, True), R("把同样的等式写进提示词，让模型自己算、自己判——检验“校验不进提示词”这条铁律", 15, INK)]])
footnote(s, "数据：CORD-v2（NAVER CLOVA，CC BY 4.0）；katanaml-org/invoices-donut-data-v1（MIT，合成发票）。预先登记记录与全部代码见仓库 README。")
notes(s, "评测上我们刻意防止“看着答案调参”：规则只在前 20 张上调，之后每一批新数据只跑一次，"
         "而且在出结果前把检验方法和预期写进仓库。最后一批 300 张是确认性检验，主检验事先定好。")

# ---------------------------------------------------------------- 7 核心结果

s = new()
header(s, "结果", "确认性检验：显著优于让模型自己核算", 7)
text(s, 0.6, 1.85, 3.6, 0.9, [[R("96.7%", 54, NAVY, True, latin=EN)]])
text(s, 0.6, 2.8, 3.6, 0.4, [[R("篡改检出率（145 / 150）", 14, MUTED)]])
text(s, 0.6, 3.4, 3.6, 0.9, [[R("3.3%", 54, NAVY, True, latin=EN)]])
text(s, 0.6, 4.35, 3.6, 0.4, [[R("误报率（5 / 150）", 14, MUTED)]])
cd = CategoryChartData()
cd.categories = ["对照 A　直接问", "对照 B　模型自己核算", "本项目"]
cd.add_series("篡改检出率", (0.747, 0.900, 0.967))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(4.4), Inches(1.7), Inches(8.3), Inches(3.0), cd)
plot = chart_style(gf.chart, legend=False)
plot.gap_width = 55
for idx, col in enumerate((GRAY2, GRAY1, NAVY)):
    pt = plot.series[0].points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
box(s, 0.6, 5.0, 12.13, 1.6)
text(s, 0.9, 5.12, 11.6, 1.4, [
    ([R("主检验（预先登记）", 15, RED, True), R("　13 : 3，精确 McNemar 检验 p = 0.021", 15, INK, True)], {"after": 2}),
    ([R("150 张篡改票逐张比较：只有本项目检出的 13 张，只有对照 B 检出的 3 张", 13, MUTED)], {"after": 6}),
    ([R("对照 A", 15, NAVY, True), R("　35 : 2，p ≈ 10⁻⁸", 15, INK)], {"after": 5}),
    [R("误报率", 15, NAVY, True), R("　本项目 3.3%，对照 B 2.7%，对照 A 1.3%——差异不显著", 15, INK)]])
footnote(s, "CORD-v2 train 划分，按标注筛出自洽的 420 张中取前 300 张，偶数号篡改。各组读 3 次、按多数判定。")
notes(s, "这是最重要的一页。事先定好的主检验是：在 150 张被改的票上，逐张比较本项目和“让模型自己核算”谁抓到了。"
         "13 比 3，p 等于 0.021。误报率我们如实报：本项目 3.3%，略高但与对照无显著差异——"
         "代码核对对任何一个误读都敏感，这是它检出率高的原因，也是误报的来源。")

# ---------------------------------------------------------------- 8 一致性

s = new()
header(s, "结果", "三批从未用过的数据，方向始终一致", 8)
cd = CategoryChartData()
cd.categories = ["CORD validation\n27 张篡改", "商业发票（合成）\n36 张篡改", "CORD train\n150 张篡改"]
cd.add_series("本项目", (25 / 27, 36 / 36, 145 / 150))
cd.add_series("对照 B　模型自己核算", (21 / 27, 34 / 36, 135 / 150))
cd.add_series("对照 A　直接问", (18 / 27, 33 / 36, 112 / 150))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.6), Inches(8.2), Inches(5.1), cd)
plot = chart_style(gf.chart)
plot.gap_width, plot.overlap = 70, -8
for ser, col in zip(plot.series, (NAVY, GRAY1, GRAY2)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
box(s, 9.1, 1.75, 3.63, 3.55)
text(s, 9.35, 1.95, 3.2, 3.3, [
    ([R("怎么读这张图", 16, NAVY, True)], {"after": 10}),
    ([R("每一批里，本项目的检出率都不低于两个对照。", 14, INK)], {"after": 10}),
    ([R("误报率：validation 与发票三组均为 0；train 为 3.3%、2.7%、1.3%。", 14, INK)], {"after": 10}),
    [R("发票是合成数据、版式规整，三组都接近满分，差距不显著——证明的是方法能换到对公单据上用。", 14, MUTED)]],
    line=1.15)
footnote(s, "三批均在出结果前冻结规则；validation 与发票为预先登记的评测，train 为预先登记的确认性检验。")
notes(s, "这张图说明结果不是某一批数据的偶然。三批数据方向一致；发票上大家都接近满分，"
         "我们不夸大，它证明的是方法能迁移到对公业务的单据上。")

# ---------------------------------------------------------------- 9 对照为什么失败

s = new()
header(s, "机理", "让模型自己核算时，它会读出“应该是”的数", 9)
image(s, ASSETS / "cord36_tax_tampered.png", 0.6, 1.8, w=4.0)
text(s, 0.6, 3.07, 4.0, 0.35, [[R("CORD 36：税额被改成 1,427（原为 1,727）", 12, MUTED)]])
cases = [("CORD 36", "税额印着 1,427", "“小计加 PB1 1,727 等于总额”",
          "把改后的数读回原值。同一张图让它只照抄，3 次识别全部判出不闭合"),
         ("train#125", "应付印着 11,000", "“应付总额 17,000……三项均成立”", "同样把改后的应付读回了原值"),
         ("train#397", "现金被改，没有找零行", "“规则 3 因无找零行而跳过”", "没有找零时付款本应等于应付——规则执行错误")]
y = 1.8
for tag, printed, said, why in cases:
    box(s, 4.95, y, 7.78, 1.2)
    text(s, 5.2, y + 0.12, 1.5, 0.95, [[R(tag, 14, NAVY, True, latin=EN)]], anchor="middle")
    text(s, 6.7, y + 0.1, 5.9, 1.0, [
        ([R("票面：", 13, MUTED), R(printed, 13, INK), R("　模型说：", 13, MUTED), R(said, 13, RED, True)], {"after": 3}),
        [R(why, 13, INK)]], anchor="middle", line=1.1)
    y += 1.38
box(s, 0.6, 3.6, 4.0, 2.1, fill=RED_BG)
text(s, 0.85, 3.75, 3.55, 1.9, [
    ([R("同一个模型、同一张图", 17, RED, True)], {"after": 8}),
    ([R("让它照抄，它读得出被改的数；", 15, INK)], {"after": 4}),
    ([R("让它核对，它朝等式去读数。", 15, INK)], {"after": 10}),
    [R("这正是“校验不进提示词”的原因。", 14, MUTED)]], line=1.15)
footnote(s, "引文为对照 B 在 CORD validation 与 CORD train 上漏检时，第 1 次识别给出的理由原文。")
notes(s, "对照 B 为什么输：不是它不会算，而是一旦知道要核对，它就倾向于把数字读成能对上的样子。"
         "CORD 36 最典型：票面印着 1,427，它说 1,727，然后宣布算术成立。")

# ---------------------------------------------------------------- 10 多数判定

s = new()
header(s, "设计", "三次识别按多数判定，不偏向“自洽”的那次", 10)
reads = [("第 1 次", "读出 8.700", "存疑", RED, RED_BG),
         ("第 2 次", "读出 8.700", "存疑", RED, RED_BG),
         ("第 3 次", "读成 8.500", "自洽", GREEN, GREEN_BG)]
for i, (t, got, v, c, bg) in enumerate(reads):
    y = 1.85 + i * 1.05
    box(s, 0.6, y, 4.2, 0.85, fill=bg)
    text(s, 0.85, y, 3.8, 0.85, [[R(t + "　", 15, NAVY, True), R(got, 15, INK), R("　→ " + v, 15, c, True)]],
         anchor="middle")
text(s, 0.6, 5.05, 4.2, 0.6, [[R("CORD 34：税额 8.500 被改成 8.700，3 次识别各自的结果", 13, MUTED)]])
for i, (rule, res, c, why) in enumerate([
        ("旧规则：自洽优先", "可信（漏检）", GREEN, "先挑能对上账的那次读数——恰好挑中了被抹平的那次"),
        ("新规则：多数判定", "存疑（检出）", RED, "每次各自判定，至少一半存疑即存疑")]):
    y = 1.85 + i * 1.6
    box(s, 5.3, y, 7.43, 1.35)
    text(s, 5.6, y + 0.15, 3.5, 0.45, [[R(rule, 17, NAVY, True)]])
    text(s, 9.2, y + 0.15, 3.3, 0.45, [[R(res, 17, c, True)]], align=PP_ALIGN.RIGHT)
    text(s, 5.6, y + 0.68, 6.9, 0.55, [[R(why, 14, INK)]])
text(s, 5.3, 5.2, 1.7, 1.0, [[R("7 : 0", 36, NAVY, True, latin=EN)]], anchor="middle")
text(s, 7.1, 5.2, 5.6, 1.0, [[R("两批未参与选规则的数据上，", 15, INK)],
                             [R("新规则多检出 7 张、少检出 0 张（p = 0.016）", 15, INK)]], anchor="middle", line=1.15)
footnote(s, "例子为 CORD test #34（开发集）。规则在开发集上选定；留出集 3 : 0、CORD validation 4 : 0。")
notes(s, "投票也有讲究。旧规则“先挑自洽的读数”看似稳妥，实际上会偏向被模型抹平的那次。"
         "改成多数判定后，在两批新数据上多抓了 7 张，没有一张反向。")

# ---------------------------------------------------------------- 11 业务场景

s = new()
header(s, "落地", "对公与供应链金融的单据预审", 11)
flow = [("单据上传", TINT, NAVY), ("自动核验", NAVY, WHITE)]
for i, (t, bg, col) in enumerate(flow):
    x = 0.6 + i * 2.35
    box(s, x, 1.9, 1.95, 0.9, fill=bg)
    text(s, x, 1.9, 1.95, 0.9, [[R(t, 16, col, True)]], anchor="middle", align=PP_ALIGN.CENTER)
    arrow(s, x + 1.98, 2.18)
outs = [("存疑", "转人工复核，附上失败的等式与差额", RED, RED_BG),
        ("无法核验", "票面缺核验所需的行，按原流程审核", MUTED, TINT),
        ("可信", "进入常规抽检", GREEN, GREEN_BG)]
for i, (t, d, c, bg) in enumerate(outs):
    y = 1.6 + i * 0.62
    box(s, 5.3, y, 7.43, 0.52, fill=bg)
    text(s, 5.5, y, 7.1, 0.52, [[R(t + "　", 15, c, True), R(d, 14, INK)]], anchor="middle")
rows = [["单据", "官方查验渠道", "本工具的角色"],
        ["境内增值税发票、数电票", "税务总局全国增值税发票查验平台", "查验前的字段录入、批量预筛"],
        ["境外发票（跨境贸易、港澳报销）", "无统一渠道", R("主要对象", 14, RED, True)],
        ["B2B 商业发票（贸易融资）", "无统一渠道", R("主要对象", 14, RED, True)],
        ["收据、小票等费用凭证", "无统一渠道", R("主要对象", 14, RED, True)]]
table(s, 0.6, 3.75, 12.13, rows, [4.2, 4.33, 3.6], size=14, row_h=0.52)
footnote(s, "官方查验平台限查税务系统开具的发票（inv-veri.chinatax.gov.cn/fpcysm.html）；本工具不替代官方查验。")
notes(s, "银行评委最可能问：增值税发票不是能去税务局查吗？能，而且应该查。我们覆盖的是官方管不到的单据："
         "境外发票、贸易融资里的商业发票、收据小票。输出直接接进人工复核队列，审核员拿到的是具体算式。")

# ---------------------------------------------------------------- 12 演示

s = new()
header(s, "演示", "每个判定都能指到具体的等式", 12)
box(s, 0.58, 1.68, 8.04, 5.04, fill=WHITE, line=LINE, radius=0.01)
image(s, ASSETS / "demo_cord36_tax.png", 0.6, 1.7, w=8.0)          # ppt/shoot_demo.py 截的真实界面，1600x1000
hl = box(s, 2.23, 3.43, 6.29, 0.46, fill=WHITE, line=RED, radius=0.12)   # 圈出不成立的校验一（截图坐标 x338–1572、y356–428）
hl.fill.background()
hl.line.width = Pt(2.25)
text(s, 0.6, 6.8, 8.0, 0.25, [[R("真实界面：CORD 36 税额被改，回放 2026-09-29 的真实识别结果。现场演示为实时识别。", 10, MUTED)]])
text(s, 9.0, 1.8, 3.73, 4.9, [
    ([R("现场怎么演示", 17, NAVY, True)], {"after": 10}),
    ([R("•  点选或上传一张单据，3 次识别并行完成", 14, INK)], {"after": 8}),
    ([R("•  判定横幅之下逐条列出等式：代入的数、票面印的数、差额", 14, INK)], {"after": 8}),
    ([R("•  单个 exe 双击即用，本机运行；网络或接口出问题时，回放之前那次真实识别的结果并标明时间", 14, INK)],
     {"after": 8}),
    [R("•  示例：港式小票、CORD 36 / 23、商业发票原图与被改版本", 14, INK)]], line=1.15)
notes(s, "现场演示流程：先点 CORD 36 原图，判可信；再点被改的那张，判存疑，并指出“小计加税不等于应付，差 300”。"
         "再点商业发票被改版本，指到第 2 行：4 × 28.08 应为 112.32，票面 117.32。"
         "若被问到港式小票为什么有一次判存疑：那次把一条折扣 6.00 读成了 5.00，只有它一次不闭合，"
         "3 次里 1 次，多数仍判可信——这正是读 3 次、按多数判定的原因。")

# ---------------------------------------------------------------- 13 边界

s = new()
header(s, "边界", "它做不到什么，我们先说", 13)
limits = [("只查票面算术", "不做像素取证。把整条依赖链成套改平的伪造检不出——本方法提高的是伪造成本。"),
          ("误报率 3.3%", "代码核对对任何误读都敏感；多数判定已把单次识别的 4.2% 压到 3.3%。"),
          ("容差半个货币单位", "88,000.00 改成 88,000.08 这类改动被视为舍入。"),
          ("实验只改一个数字", "真实造假可能改多处；多处改动需要同时骗过多条等式（见第 5 页）。"),
          ("发票为合成数据", "版式规整、印刷清晰；真实拍照的单据更难读。"),
          ("只测了一个视觉模型", "DeepSeek；换模型需要重新评测。")]
for i, (t, d) in enumerate(limits):
    x = 0.6 + (i % 3) * 4.1
    y = 1.85 + (i // 3) * 2.25
    box(s, x, y, 3.8, 1.95)
    text(s, x + 0.3, y + 0.25, 3.25, 0.45, [[R(t, 18, NAVY, True)]])
    text(s, x + 0.3, y + 0.85, 3.25, 1.2, [[R(d, 14, INK)]], line=1.2)
notes(s, "局限主动说：只查算术不做图像取证；误报 3.3%；容差会放过小于半个单位的改动；"
         "实验只改一个数；发票是合成的；只测了一个模型。每一条都写在仓库里。")

# ---------------------------------------------------------------- 14 下一步

s = new()
header(s, "下一步", "从公开数据走向真实业务", 14)
nexts = [("中国金融单据", "支票、收据、发票上的“大写金额 = 小写金额”——又一个免费校验位，而且大写是文字不是数字"),
         ("按行转写", "裁出单行再读，对付“整张一起读时被抹平”的稳定误读"),
         ("多模型交叉", "不同模型的误读若互不相关，交叉读数可进一步降低漏检（待验证）"),
         ("接入复核流程", "存疑单据附算式进入人工复核队列；记录处置结果，持续评估误报与漏检")]
for i, (t, d) in enumerate(nexts):
    y = 1.8 + i * 1.22
    circle_num(s, 0.6, y + 0.12, i + 1, d=0.56, size=17)
    text(s, 1.45, y, 3.2, 0.8, [[R(t, 19, NAVY, True)]], anchor="middle")
    text(s, 4.7, y, 8.0, 0.8, [[R(d, 15, INK)]], anchor="middle", line=1.15)
notes(s, "下一步四件事，第一件最贴近工行的业务：中国的金融单据普遍有大写金额，这是天然的第二份记录。")

# ---------------------------------------------------------------- 15 结尾

s = new(dark=True)
text(s, 0.8, 2.0, 11.7, 1.2, [[R("票面会说话", 50, WHITE, True)]])
text(s, 0.8, 3.2, 11.7, 0.8, [[R("前提是让模型只负责读，让代码负责算。", 26, ICE)]])
text(s, 0.8, 4.7, 11.7, 1.0, [
    ([R("代码、数据处理、全部实验与预先登记记录公开：", 15, ICE)], {"after": 4}),
    [R("github.com/baimingyang98/receipt-integrity", 17, WHITE, True, latin=EN)]])
text(s, 0.8, 6.3, 11.7, 0.5, [[R("谢谢评委老师", 18, WHITE)]])
notes(s, "收尾：票面自己会说话，前提是让模型只负责读、让代码负责算。所有代码和实验记录都公开，欢迎提问。")

prs.save(OUT)
print("已生成:", OUT, "-", len(prs.slides), "页")
